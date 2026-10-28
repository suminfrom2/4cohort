import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from io import BytesIO
from openpyxl.styles import PatternFill
import os

# 페이지 기본 설정
st.set_page_config(
    page_title="RWD 코호트 분석 웹 대시보드",
    page_icon="[RWD]",
    layout="wide"
)

# 한글 폰트 설정
plt.rc('font', family='NanumGothic')
plt.rcParams['axes.unicode_minus'] = False

st.title("RWD 기반 2단계 하위 집단 필터링 및 엑셀 리스트업")
st.markdown("본 시스템은 기저변수(T1, T3)를 활용해 원하는 특정 조건(예: BMI >= 25)으로 코호트를 걸러내고, 그 안에서 두 번째 변수를 기준으로 두 개의 군으로 분할하여 약물 처방 현황과 전체 임상 변수를 분석합니다.")

# 1. 파일 업로드 기능
st.sidebar.header("1. 데이터 파일 업로드")
uploaded_file = st.sidebar.file_uploader("실습데이터셋 엑셀 파일을 선택하세요.", type=["xlsx"])

if uploaded_file is not None:
    try:
        # 데이터 캐싱하여 속도 개선
        @st.cache_data
        def load_data(file):
            t1_df = pd.read_excel(file, sheet_name='T1_대상자')
            t3_df = pd.read_excel(file, sheet_name='T3_기저특성')
            t2_df = pd.read_excel(file, sheet_name='T2_처방')
            mapping_df = pd.read_excel(file, sheet_name='성분매핑')
            return t1_df, t3_df, t2_df, mapping_df
        
        t1_df, t3_df, t2_df, mapping_df = load_data(uploaded_file)
        
        # 데이터 통합
        if 'AGE' in t1_df.columns:
            t1_df = t1_df.drop(columns=['AGE'])
            
        t2_mapped = pd.merge(t2_df, mapping_df, on='DRUG_NAME', how='left')
        patient_drugs = t2_mapped.groupby('PT_ID')['DRUG_CLASS'].apply(lambda x: list(set(x.dropna()))).reset_index()
        patient_drugs.columns = ['PT_ID', 'PRESCRIBED_CLASSES']
        
        full_master = pd.merge(t1_df, t3_df.drop(columns=['INDEX_DATE'], errors='ignore'), on='PT_ID', how='outer')
        full_master = pd.merge(full_master, patient_drugs, on='PT_ID', how='left')
        full_master['PRESCRIBED_CLASSES'] = full_master['PRESCRIBED_CLASSES'].apply(lambda x: x if isinstance(x, list) else [])
        
        exclude_cols = ['PT_ID', 'PRESCRIBED_CLASSES', 'INDEX_DATE']
        analysis_vars = [col for col in full_master.columns if col not in exclude_cols]
        
        # 대시보드 위젯 설정
        st.sidebar.markdown("---")
        st.sidebar.header("2. 조건 설정")
        
        # Section 1 구성
        st.sidebar.subheader("[Section 1: 1차 코호트 필터]")
        v1 = st.sidebar.selectbox("1차 필터 변수", options=analysis_vars, index=analysis_vars.index('BMI') if 'BMI' in analysis_vars else 0)
        op1 = st.sidebar.selectbox("조건 식", options=['>= (이상)', '<= (이하)', '== (일치)'], index=0, key='op1')
        val1_input = st.sidebar.text_input("기준값 입력", value='25', key='val1')
        
        # Section 2 구성
        st.sidebar.subheader("[Section 2: 2차 세부 분석 그룹 분류]")
        v2 = st.sidebar.selectbox("2차 분류 변수", options=analysis_vars, index=analysis_vars.index('FBG_BL') if 'FBG_BL' in analysis_vars else 0)
        op2 = st.sidebar.selectbox("분류 기준식", options=['>= (이상)', '<= (이하)', '== (일치)'], index=0, key='op2')
        val2_input = st.sidebar.text_input("분류 기준값", value='100', key='val2')
        
        # 값 변환
        def parse_val(v, df_series):
            try:
                if df_series.dtype in [np.float64, np.int64]:
                    return float(v)
                return str(v)
            except:
                return v
        
        val1 = parse_val(val1_input, full_master[v1])
        val2 = parse_val(val2_input, full_master[v2])
        
        # --- [1단계] 1차 필터링 적용 ---
        cohort_df = full_master.copy()
        if op1 == '>= (이상)':
            cohort_df = cohort_df[cohort_df[v1] >= val1]
        elif op1 == '<= (이하)':
            cohort_df = cohort_df[cohort_df[v1] <= val1]
        else:
            cohort_df = cohort_df[cohort_df[v1].astype(str) == str(val1)]
            
        # UI 레이아웃 화면 표시
        col1, col2 = st.columns(2)
        
        with col1:
            st.info(f"* [Section 1] 1차 필터링 코호트 결과\n* 필터 조건: {v1} {op1} {val1}\n* 추출된 코호트 대상자: {len(cohort_df)}명 / 전체 {len(full_master)}명")
            
        if len(cohort_df) == 0:
            st.warning("1차 필터 조건에 만족하는 환자가 없습니다. 왼쪽 사이드바의 입력값을 변경해주세요.")
        else:
            # --- [2단계] 2차 그룹 분할 ---
            if op2 == '>= (이상)':
                mask = cohort_df[v2] >= val2
                g_a_desc = f"1차 기준({v1} {op1} {val1})을 충족하면서, 2차 기준({v2}가 {val2} 이상)인 집단"
                g_b_desc = f"1차 기준({v1} {op1} {val1})을 충족하지만, 2차 기준({v2}가 {val2} 미만)인 집단"
            elif op2 == '<= (이하)':
                mask = cohort_df[v2] <= val2
                g_a_desc = f"1차 기준({v1} {op1} {val1})을 충족하면서, 2차 기준({v2}가 {val2} 이하)인 집단"
                g_b_desc = f"1차 기준({v1} {op1} {val1})을 충족하지만, 2차 기준({v2}가 {val2} 초과)인 집단"
            else:
                mask = cohort_df[v2].astype(str) == str(val2)
                g_a_desc = f"1차 기준({v1} {op1} {val1})을 충족하면서, 2차 기준({v2}가 {val2}와 일치)인 집단"
                g_b_desc = f"1차 기준({v1} {op1} {val1})을 충족하지만, 2차 기준({v2}가 {val2}와 불일치)인 집단"
                
            cohort_df['Analysis_Group'] = np.where(mask, 'Group A', 'Group B')
            group_a = cohort_df[cohort_df['Analysis_Group'] == 'Group A']
            group_b = cohort_df[cohort_df['Analysis_Group'] == 'Group B']
            
            pct_a = (len(group_a)/len(cohort_df)*100)
            pct_b = (len(group_b)/len(cohort_df)*100)
            
            with col2:
                st.warning(f"* [Section 2] 그룹별 총체적 임상 특성 정의\n* Group A 정의: {g_a_desc} (총 {len(group_a)}명, {pct_a:.1f}%)\n* Group B 정의: {g_b_desc} (총 {len(group_b)}명, {pct_b:.1f}%)")
            
            # --- [3단계] 최종 그룹별 약물 처방 현황 분석 및 시각화 ---
            st.write("### 두 그룹의 고혈압 처방 약물 계열(Class) 분포 비교")
            
            def get_drug_counts(df_sub):
                classes_list = [c for sublist in df_sub['PRESCRIBED_CLASSES'] for c in sublist]
                s = pd.Series(classes_list).value_counts()
                return s
                
            drug_a = get_drug_counts(group_a).to_frame(name='Group A (명)')
            drug_b = get_drug_counts(group_b).to_frame(name='Group B (명)')
            drug_comparison = pd.concat([drug_a, drug_b], axis=1).fillna(0).astype(int)
            
            if not drug_comparison.empty:
                fig, ax = plt.subplots(figsize=(10, 4.5))
                drug_comparison.plot(kind='bar', ax=ax, color=['#3498db', '#e74c3c'], edgecolor='black', zorder=3)
                
                for p in ax.patches:
                    height = int(p.get_height())
                    if height > 0:
                        ax.annotate(f'{height}명',
                                    (p.get_x() + p.get_width() / 2., height),
                                    ha='center', va='center',
                                    xytext=(0, 6),
                                    textcoords='offset points',
                                    fontsize=9, fontweight='bold')
                                    
                plt.title('세부 비교 그룹별 처방 약물계열 현황', fontsize=12, pad=10)
                plt.ylabel('처방 환자 수 (명)', fontsize=10)
                plt.xlabel('약물 계열 (Drug Class)', fontsize=10)
                plt.xticks(rotation=0, ha='center', fontsize=9)
                plt.ylim(0, drug_comparison.values.max() * 1.15)
                plt.grid(axis='y', linestyle='--', alpha=0.5, zorder=0)
                plt.tight_layout()
                st.pyplot(fig)
                
                # 표 출력
                summary_table = drug_comparison.copy()
                summary_table['Group A 비율 (%)'] = (summary_table['Group A (명)'] / len(group_a) * 100).round(1) if len(group_a) > 0 else 0
                summary_table['Group B 비율 (%)'] = (summary_table['Group B (명)'] / len(group_b) * 100).round(1) if len(group_b) > 0 else 0
                summary_table = summary_table[['Group A (명)', 'Group A 비율 (%)', 'Group B (명)', 'Group B 비율 (%)']]
                st.write("#### 약물 처방 현황 상세 분석표")
                st.dataframe(summary_table, use_container_width=True)
            else:
                st.write("해당 하위 코호트 내 처방 약물 데이터가 존재하지 않습니다.")
                
            # --- [4단계] 통합 리스트업 및 개별 시트 엑셀 버퍼 다운로드 ---
            st.write("### 전체 변수가 포함된 통합 환자 리스트업 (상위 10명 미리보기)")
            display_cols = ['PT_ID', 'Analysis_Group', v1, v2] + [col for col in analysis_vars if col not in [v1, v2]] + ['PRESCRIBED_CLASSES']
            preview_df = cohort_df[display_cols].copy()
            preview_df['PRESCRIBED_CLASSES'] = preview_df['PRESCRIBED_CLASSES'].apply(lambda x: ", ".join(x) if x else "없음")
            st.dataframe(preview_df.head(10), use_container_width=True)
            
            # 중복 약물 복용 여부 헬퍼
            def is_duplicate_user(classes_list):
                if isinstance(classes_list, list):
                    return 'ACE 억제제' in classes_list and '티아지드계 이뇨제' in classes_list
                return False
                
            # 엑셀 다운로드 파일 생성 (In-memory BytesIO 사용)
            excel_buffer = BytesIO()
            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                for grp_name, grp_df in [('Group_A', group_a), ('Group_B', group_b)]:
                    sheet_df = grp_df[display_cols].copy()
                    
                    is_dup_mask = sheet_df['PRESCRIBED_CLASSES'].apply(is_duplicate_user)
                    dup_pt_ids = sheet_df.loc[is_dup_mask, 'PT_ID'].tolist()
                    
                    sheet_df['PRESCRIBED_CLASSES'] = sheet_df['PRESCRIBED_CLASSES'].apply(lambda x: ", ".join(x) if x else "없음")
                    sheet_df.to_excel(writer, sheet_name=grp_name, index=False)
                    
                    # 스타일 적용
                    worksheet = writer.sheets[grp_name]
                    yellow_fill = PatternFill(start_color="FFFFD1", end_color="FFFFD1", fill_type="solid")
                    
                    for row_idx, val in enumerate(is_dup_mask, start=2):
                        if val:
                            for col_idx in range(1, len(display_cols) + 1):
                                cell = worksheet.cell(row=row_idx, column=col_idx)
                                cell.fill = yellow_fill
                                
                    dup_col_start = len(display_cols) + 2
                    worksheet.cell(row=1, column=dup_col_start, value="중복 복용 환자 ID 목록")
                    for i, pt_id in enumerate(dup_pt_ids, start=2):
                        worksheet.cell(row=i, column=dup_col_start, value=pt_id)
            
            excel_buffer.seek(0)
            
            st.markdown("---")
            st.success("다운로드할 분석 결과 엑셀 파일이 서버 메모리에 빌드되었습니다!")
            st.download_button(
                label="맞춤형 다중 시트 분석 결과 엑셀 다운로드",
                data=excel_buffer,
                file_name="cohort_analysis_result_multisheet.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            
    except Exception as e:
        st.error(f"데이터를 처리하는 도중 오류가 발생했습니다: {e}")

else:
    st.info("왼쪽 사이드바에서 실습데이터셋 엑셀 파일을 업로드해 주세요.")
